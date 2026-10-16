from django.test import TestCase
from django.urls import reverse
from django.contrib.auth.models import User

class DashboardTestCase(TestCase):
    def setUp(self):
        # Create non-staff student user
        self.student_user = User.objects.create_user(
            username='student@test.com',
            email='student@test.com',
            password='Password123!',
            is_staff=False
        )
        # Create staff admin user
        self.staff_user = User.objects.create_user(
            username='admin@test.com',
            email='admin@test.com',
            password='Password123!',
            is_staff=True
        )

    def test_unauthenticated_request_redirects_to_login(self):
        url = reverse('dashboard:index')
        response = self.client.get(url)
        self.assertRedirects(response, reverse('dashboard:login'))

    def test_non_staff_student_is_rejected(self):
        # Log in as non-staff student
        self.client.login(username='student@test.com', password='Password123!')
        url = reverse('dashboard:index')
        response = self.client.get(url)
        # Staff required decorator redirects non-staff user to dashboard:login with error message
        self.assertRedirects(response, reverse('dashboard:login'))

    def test_staff_admin_can_access_dashboard(self):
        self.client.login(username='admin@test.com', password='Password123!')
        url = reverse('dashboard:index')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'dashboard/index.html')


class AdminAccountSettingsTestCase(TestCase):
    def setUp(self):
        self.staff_admin = User.objects.create_user(
            username='admin_boss',
            email='admin@cdac.test',
            password='CurrentStrongPassword123!',
            is_staff=True
        )
        self.other_user = User.objects.create_user(
            username='existing_user',
            email='existing@cdac.test',
            password='SomePassword123!',
            is_staff=False
        )

    def test_unauthenticated_access_redirects(self):
        for route_name in ['account_settings', 'change_username', 'change_password']:
            url = reverse(f'dashboard:{route_name}')
            response = self.client.get(url)
            self.assertRedirects(response, reverse('dashboard:login'))

    def test_non_staff_access_rejected(self):
        self.client.login(username='existing_user', password='SomePassword123!')
        url = reverse('dashboard:account_settings')
        response = self.client.get(url)
        self.assertRedirects(response, reverse('dashboard:login'))

    def test_staff_can_view_account_settings_page(self):
        self.client.login(username='admin_boss', password='CurrentStrongPassword123!')
        url = reverse('dashboard:account_settings')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'dashboard/account_settings.html')
        self.assertContains(response, 'admin_boss')
        self.assertContains(response, 'Change Username')
        self.assertContains(response, 'Change Password')

    def test_change_username_success(self):
        self.client.login(username='admin_boss', password='CurrentStrongPassword123!')
        url = reverse('dashboard:change_username')
        response = self.client.post(url, {
            'action': 'change_username',
            'new_username': 'super_admin_new',
            'current_password': 'CurrentStrongPassword123!'
        })
        self.assertRedirects(response, reverse('dashboard:account_settings'))

        # Verify DB updated
        self.staff_admin.refresh_from_db()
        self.assertEqual(self.staff_admin.username, 'super_admin_new')

        # Verify can log in with new username
        self.client.logout()
        login_success = self.client.login(username='super_admin_new', password='CurrentStrongPassword123!')
        self.assertTrue(login_success)

    def test_change_username_wrong_password(self):
        self.client.login(username='admin_boss', password='CurrentStrongPassword123!')
        url = reverse('dashboard:change_username')
        response = self.client.post(url, {
            'action': 'change_username',
            'new_username': 'new_boss',
            'current_password': 'WrongPassword123!'
        })
        self.assertEqual(response.status_code, 200)
        self.assertFormError(response.context['username_form'], 'current_password', 'Incorrect current password.')
        self.staff_admin.refresh_from_db()
        self.assertEqual(self.staff_admin.username, 'admin_boss')

    def test_change_username_duplicate_rejected(self):
        self.client.login(username='admin_boss', password='CurrentStrongPassword123!')
        url = reverse('dashboard:change_username')
        response = self.client.post(url, {
            'action': 'change_username',
            'new_username': 'existing_user',
            'current_password': 'CurrentStrongPassword123!'
        })
        self.assertEqual(response.status_code, 200)
        self.assertFormError(response.context['username_form'], 'new_username', 'This username is already taken. Please choose another.')
        self.staff_admin.refresh_from_db()
        self.assertEqual(self.staff_admin.username, 'admin_boss')

    def test_change_username_same_as_current_rejected(self):
        self.client.login(username='admin_boss', password='CurrentStrongPassword123!')
        url = reverse('dashboard:change_username')
        response = self.client.post(url, {
            'action': 'change_username',
            'new_username': 'admin_boss',
            'current_password': 'CurrentStrongPassword123!'
        })
        self.assertEqual(response.status_code, 200)
        self.assertFormError(response.context['username_form'], 'new_username', 'New username must be different from your current username.')

    def test_change_password_success_and_session_preserved(self):
        self.client.login(username='admin_boss', password='CurrentStrongPassword123!')
        url = reverse('dashboard:change_password')
        response = self.client.post(url, {
            'action': 'change_password',
            'current_password': 'CurrentStrongPassword123!',
            'new_password': 'BrandNewSecurePassword456!',
            'confirm_password': 'BrandNewSecurePassword456!'
        })
        self.assertRedirects(response, reverse('dashboard:account_settings'))

        # Check session is still authenticated without needing to re-login
        dashboard_url = reverse('dashboard:index')
        dash_response = self.client.get(dashboard_url)
        self.assertEqual(dash_response.status_code, 200)

        # Verify old password no longer works
        self.client.logout()
        old_login_success = self.client.login(username='admin_boss', password='CurrentStrongPassword123!')
        self.assertFalse(old_login_success)

        # Verify new password works
        new_login_success = self.client.login(username='admin_boss', password='BrandNewSecurePassword456!')
        self.assertTrue(new_login_success)

    def test_change_password_wrong_current_password(self):
        self.client.login(username='admin_boss', password='CurrentStrongPassword123!')
        url = reverse('dashboard:change_password')
        response = self.client.post(url, {
            'action': 'change_password',
            'current_password': 'WrongPassword!',
            'new_password': 'BrandNewSecurePassword456!',
            'confirm_password': 'BrandNewSecurePassword456!'
        })
        self.assertEqual(response.status_code, 200)
        self.assertFormError(response.context['password_form'], 'current_password', 'Incorrect current password.')

    def test_change_password_mismatch(self):
        self.client.login(username='admin_boss', password='CurrentStrongPassword123!')
        url = reverse('dashboard:change_password')
        response = self.client.post(url, {
            'action': 'change_password',
            'current_password': 'CurrentStrongPassword123!',
            'new_password': 'BrandNewSecurePassword456!',
            'confirm_password': 'DifferentPassword456!'
        })
        self.assertEqual(response.status_code, 200)
        self.assertFormError(response.context['password_form'], 'confirm_password', "The two password fields didn't match.")

    def test_change_password_too_short_validation(self):
        self.client.login(username='admin_boss', password='CurrentStrongPassword123!')
        url = reverse('dashboard:change_password')
        response = self.client.post(url, {
            'action': 'change_password',
            'current_password': 'CurrentStrongPassword123!',
            'new_password': '123',
            'confirm_password': '123'
        })
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['password_form'].has_error('new_password'))


import time
from django.conf import settings

class AdminAutoLogoutSecurityTestCase(TestCase):
    def setUp(self):
        self.admin_user = User.objects.create_user(
            username='admin_security_test',
            email='security@cdac.test',
            password='StrongSecurityPassword123!',
            is_staff=True
        )

    def test_security_settings_configured(self):
        self.assertTrue(settings.SESSION_EXPIRE_AT_BROWSER_CLOSE)
        self.assertEqual(settings.SESSION_COOKIE_AGE, 3600)
        self.assertFalse(settings.SESSION_SAVE_EVERY_REQUEST)
        self.assertTrue(settings.SESSION_COOKIE_HTTPONLY)
        self.assertEqual(settings.SESSION_COOKIE_SAMESITE, 'Lax')

    def test_no_cache_headers_on_admin_pages(self):
        self.client.login(username='admin_security_test', password='StrongSecurityPassword123!')
        response = self.client.get(reverse('dashboard:index'))
        self.assertEqual(response.status_code, 200)
        self.assertIn('no-store', response.headers.get('Cache-Control', ''))
        self.assertIn('no-cache', response.headers.get('Cache-Control', ''))
        self.assertEqual(response.headers.get('Pragma'), 'no-cache')
        self.assertEqual(response.headers.get('Expires'), '0')

    def test_no_cache_headers_not_on_public_pages(self):
        response = self.client.get(reverse('core:home'))
        self.assertEqual(response.status_code, 200)
        cache_control = response.headers.get('Cache-Control', '')
        self.assertNotIn('no-store', cache_control)

    def test_beacon_logout_endpoint(self):
        self.client.login(username='admin_security_test', password='StrongSecurityPassword123!')
        url = reverse('dashboard:beacon_logout')
        response = self.client.post(url)
        self.assertEqual(response.status_code, 204)

        # Subsequent dashboard request must be redirected to login
        dash_response = self.client.get(reverse('dashboard:index'))
        self.assertRedirects(dash_response, reverse('dashboard:login'))

    def test_absolute_one_hour_timeout(self):
        self.client.login(username='admin_security_test', password='StrongSecurityPassword123!')
        
        # Set session admin_login_time to 3601 seconds ago (exceeded 1 hour)
        session = self.client.session
        session['admin_login_time'] = time.time() - 3601
        session.save()

        # Next request must be logged out and redirected to login
        response = self.client.get(reverse('dashboard:index'))
        self.assertRedirects(response, reverse('dashboard:login'))

        # Check session was invalidated
        dash_response = self.client.get(reverse('dashboard:index'))
        self.assertRedirects(dash_response, reverse('dashboard:login'))

    def test_active_session_under_one_hour_not_logged_out(self):
        self.client.login(username='admin_security_test', password='StrongSecurityPassword123!')
        
        # Set session admin_login_time to 1800 seconds ago (30 minutes)
        session = self.client.session
        session['admin_login_time'] = time.time() - 1800
        session.save()

        # Request succeeds
        response = self.client.get(reverse('dashboard:index'))
        self.assertEqual(response.status_code, 200)


import io
import datetime
import openpyxl
from courses.models import Course
from certificates.models import Certificate

class CertificateExcelExportTestCase(TestCase):
    def setUp(self):
        self.staff_admin = User.objects.create_user(
            username='admin_excel_tester',
            email='exceladmin@cdac.test',
            password='StrongPassword123!',
            is_staff=True
        )
        self.student = User.objects.create_user(
            username='student_cert_user',
            email='student_cert@cdac.test',
            password='StudentPass123!',
            is_staff=False
        )
        self.course = Course.objects.create(
            name="Python Web Frameworks",
            category="web-development",
            duration="3 Months",
            description="Django full stack development"
        )
        # Create certificates with distinct dates
        self.cert1 = Certificate.objects.create(
            certificate_id="RAS-2026-PY-001",
            first_name="Alice",
            last_name="Smith",
            student_name="Alice Smith",
            course=self.course,
            student=self.student,
            course_start_date=datetime.date(2026, 1, 10),
            course_end_date=datetime.date(2026, 4, 10),
            issue_date=datetime.date(2026, 5, 1),
            grade="A+"
        )
        self.cert2 = Certificate.objects.create(
            certificate_id="RAS-2026-PY-002",
            first_name="Bob",
            last_name="Johnson",
            student_name="Bob Johnson",
            course=self.course,
            issue_date=datetime.date(2026, 5, 15),
            grade="A"
        )
        self.cert_outside = Certificate.objects.create(
            certificate_id="RAS-2026-PY-003",
            first_name="Charlie",
            last_name="Brown",
            student_name="Charlie Brown",
            course=self.course,
            issue_date=datetime.date(2026, 8, 20),
            grade="B+"
        )

    def test_unauthenticated_or_non_staff_export_blocked(self):
        url = reverse('dashboard:export_certificates_excel')
        # Unauthenticated
        response = self.client.get(url, {'start_date': '2026-05-01', 'end_date': '2026-05-31'})
        self.assertRedirects(response, reverse('dashboard:login'))

        # Non-staff student
        self.client.login(username='student_cert_user', password='StudentPass123!')
        student_response = self.client.get(url, {'start_date': '2026-05-01', 'end_date': '2026-05-31'})
        self.assertRedirects(student_response, reverse('dashboard:login'))

    def test_export_valid_date_range_downloads_xlsx(self):
        self.client.login(username='admin_excel_tester', password='StrongPassword123!')
        url = reverse('dashboard:export_certificates_excel')
        response = self.client.get(url, {'start_date': '2026-05-01', 'end_date': '2026-05-20'})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.headers.get('Content-Type'),
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('certificates_01-05-2026_to_20-05-2026.xlsx', response.headers.get('Content-Disposition', ''))

        # Inspect workbook content
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.title, "Certificates")

        # Check headers are present and bold
        header_vals = [cell.value for cell in ws[1]]
        self.assertIn("Certificate ID", header_vals)
        self.assertIn("Student Name", header_vals)
        self.assertIn("Course Name", header_vals)
        self.assertIn("Issue Date", header_vals)
        self.assertTrue(ws.cell(row=1, column=1).font.bold)

        # cert1 and cert2 should be included (rows 2 and 3), cert_outside should NOT be included
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(data_rows), 2)

        cert_ids = [row[0] for row in data_rows]
        self.assertIn("RAS-2026-PY-001", cert_ids)
        self.assertIn("RAS-2026-PY-002", cert_ids)
        self.assertNotIn("RAS-2026-PY-003", cert_ids)

        # Verify date format DD-MM-YYYY
        row1 = data_rows[0]
        issue_date_idx = header_vals.index("Issue Date")
        self.assertEqual(row1[issue_date_idx], "01-05-2026")

    def test_export_start_date_after_end_date_rejected(self):
        self.client.login(username='admin_excel_tester', password='StrongPassword123!')
        url = reverse('dashboard:export_certificates_excel')
        response = self.client.get(url, {'start_date': '2026-05-20', 'end_date': '2026-05-01'})

        # Redirects back with error message
        self.assertRedirects(response, reverse('dashboard:manage_certificates'))

    def test_export_no_records_in_range_shows_message(self):
        self.client.login(username='admin_excel_tester', password='StrongPassword123!')
        url = reverse('dashboard:export_certificates_excel')
        response = self.client.get(url, {'start_date': '2020-01-01', 'end_date': '2020-01-31'})

        # Redirects with friendly error message rather than blank file
        self.assertRedirects(response, reverse('dashboard:manage_certificates'))

    def test_export_missing_dates_rejected(self):
        self.client.login(username='admin_excel_tester', password='StrongPassword123!')
        url = reverse('dashboard:export_certificates_excel')

        # No dates provided
        response = self.client.get(url)
        self.assertRedirects(response, reverse('dashboard:manage_certificates'))

        # Only one date provided
        response2 = self.client.get(url, {'start_date': '2026-05-01'})
        self.assertRedirects(response2, reverse('dashboard:manage_certificates'))

    def test_formula_injection_sanitization(self):
        self.client.login(username='admin_excel_tester', password='StrongPassword123!')
        # Create certificate with potential formula strings
        dangerous_cert = Certificate.objects.create(
            certificate_id="RAS-2026-INJ-001",
            first_name="=cmd",
            last_name="|' /C calc'!A0",
            student_name="=cmd|' /C calc'!A0",
            course=self.course,
            issue_date=datetime.date(2026, 5, 5),
            grade="+A"
        )
        url = reverse('dashboard:export_certificates_excel')
        response = self.client.get(url, {'start_date': '2026-05-04', 'end_date': '2026-05-06'})
        self.assertEqual(response.status_code, 200)

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        student_names = [row[1] for row in data_rows]

        # Value should be prefixed with apostrophe so spreadsheet engines treat it strictly as text
        self.assertTrue(any(name.startswith("'=") for name in student_names if name))



